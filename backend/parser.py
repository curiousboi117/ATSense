import io
import os
import zipfile
import fitz  # PyMuPDF
import docx
from typing import Dict, Any

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MiB

PDF_SIGNATURE = b"%PDF"
ZIP_SIGNATURE = b"PK"


def parse_file(file_content: bytes, filename: str) -> Dict[str, Any]:
    """
    Parse a PDF or DOCX file after validating its size, extension,
    and underlying file signature.
    """
    if len(file_content) > MAX_FILE_SIZE:
        raise ValueError("File size exceeds the maximum limit of 5MB.")

    if not file_content:
        raise ValueError("The uploaded file is empty.")

    ext = os.path.splitext(filename.lower())[1]

    if ext == ".pdf":
        if not file_content.startswith(PDF_SIGNATURE):
            raise ValueError("The uploaded file is not a valid PDF.")
        return parse_pdf(file_content)

    if ext == ".docx":
        if not file_content.startswith(ZIP_SIGNATURE):
            raise ValueError("The uploaded file is not a valid DOCX file.")

        if not is_valid_docx(file_content):
            raise ValueError("The uploaded file is not a valid DOCX file.")

        return parse_docx(file_content)

    raise ValueError(
        "Unsupported file format. Only PDF and DOCX files are allowed."
    )


def is_valid_docx(file_content: bytes) -> bool:
    """
    Verify that the file is a valid ZIP archive containing
    the core DOCX structure.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(file_content)) as archive:
            names = set(archive.namelist())

            required_files = {
                "[Content_Types].xml",
                "word/document.xml",
            }

            return required_files.issubset(names)

    except (zipfile.BadZipFile, OSError):
        return False


def parse_pdf(file_content: bytes) -> Dict[str, Any]:
    """
    Extract text from PDF file content using PyMuPDF.
    """
    doc = None

    try:
        doc = fitz.open(stream=file_content, filetype="pdf")

        text_parts = []

        for page in doc:
            page_text = page.get_text()

            if page_text:
                text_parts.append(page_text)

        text = "\n".join(text_parts)

        if not text.strip():
            raise ValueError(
                "No text could be extracted from this PDF. "
                "It might be scanned or empty."
            )

        return {
            "text": text,
            "page_count": len(doc),
            "metadata": doc.metadata or {},
        }

    except ValueError:
        raise

    except Exception as e:
        raise ValueError(f"Failed to parse PDF file: {str(e)}")

    finally:
        if doc is not None:
            doc.close()


def parse_docx(file_content: bytes) -> Dict[str, Any]:
    """
    Extract text from DOCX file content using python-docx.
    """
    try:
        doc_stream = io.BytesIO(file_content)
        doc = docx.Document(doc_stream)

        text_parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Also extract text from tables.
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        text = "\n".join(text_parts)

        if not text.strip():
            raise ValueError("No text could be extracted from this DOCX file.")

        metadata = {}

        try:
            metadata = {
                "author": doc.core_properties.author or "",
                "title": doc.core_properties.title or "",
                "created": (
                    str(doc.core_properties.created)
                    if doc.core_properties.created
                    else ""
                ),
            }
        except Exception:
            pass

        return {
            "text": text,
            "page_count": None,
            "metadata": metadata,
        }

    except ValueError:
        raise

    except Exception as e:
        raise ValueError(f"Failed to parse DOCX file: {str(e)}")